"""Review annual regional Rosstat wages without assuming preliminary data are final.

The workbook is Rosstat's annual table used in the separate public reference.
This job only proposes a data change. Final status and the actual publication
month must be confirmed before a year can be used for paid calculations.
"""
import argparse
import datetime as dt
import hashlib
import io
import json
import re
import ssl
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "rosstat-wages.json"
SOURCE_PATTERN = "https://www.rosstat.gov.ru/storage/mediabank/tab4-zpl_{year}.xlsx"
REGION_START = re.compile(r"^(?:Республика |Край |Область |Автономн|г\. |Город |Москва$|Санкт-Петербург$|Севастополь$)", re.I)

# Rosstat currently omits an intermediate certificate. These pins are checked
# before extending trust and cause a visible failure if the chain changes.
ROOT_CA = "https://gu-st.ru/content/Other/doc/russiantrustedca.pem"
ROOT_FINGERPRINT = "d26d2d0231b7c39f92cc738512ba54103519e4405d68b5bd703e9788ca8ecf31"
SUB_FINGERPRINT = "bbbde2103e790b999ec62bd03cf625a5a2e7c316e10afe6a490eedead8b3fd9b"
CURRENT_SUB = "http://nuc-cdp.digital.gov.ru/cdp/subca_ssl_rsa2024.crt"
CURRENT_SUB_FINGERPRINT = "2155785036c900dbb5f1bb2a1569c80c55595bd6bf94867a29bbddbc7d88a3f2"


def number(value):
    if isinstance(value, (int, float)) and value > 1000:
        return round(float(value), 2)
    if isinstance(value, str):
        cleaned = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if re.fullmatch(r"\d{4,6}(?:\.\d{1,2})?", cleaned):
            return round(float(cleaned), 2)
    return None


def region(value):
    if not isinstance(value, str):
        return None
    name = re.sub(r"\s+", " ", value).strip().replace("Ё", "Е")
    if "Ненецкий авт" in name:
        name = re.sub(r"^в том числе\s+", "", name, flags=re.I)
    accepted = REGION_START.search(name) or name.endswith((" область", " край", " Республика", " автономный округ", " авт.округ", " авт.область")) or any(token in name for token in ("Санкт-Петербург", "Ненецкий", "Ханты-Мансийский"))
    if accepted and not any(marker in name.lower() for marker in ("российская федерация", "федеральный округ", "в том числе", "районы крайнего")):
        return name
    return None


def extract(book):
    years, preliminary = {}, set()
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        notes = {match.group(1) for row in rows for value in row if isinstance(value, str)
                 if (match := re.search(r"(?:^|\n)\s*(\d+)\)\s*Предварительные данные", value, re.I))}
        for index, row in enumerate(rows):
            columns = {}
            for col, value in enumerate(row):
                label = str(value).strip()
                match = re.match(r"^(202[4-9])", label)
                if match:
                    columns[col] = int(match.group(1))
                    if any(f"{note})" in label for note in notes):
                        preliminary.add(int(match.group(1)))
            if len(columns) < 2 or region(row[0]):
                continue
            for values in rows[index + 1:]:
                label = next((region(item) for item in values[:4] if region(item)), None)
                if not label:
                    continue
                for col, year in columns.items():
                    if col < len(values) and (salary := number(values[col])):
                        years.setdefault(year, {})[label] = salary
    return years, preliminary


def verified_rosstat_context():
    bootstrap = ssl._create_unverified_context()
    with urllib.request.urlopen(ROOT_CA, timeout=30, context=bootstrap) as response:
        certificates = re.findall(rb"-----BEGIN CERTIFICATE-----.*?-----END CERTIFICATE-----", response.read(), re.S)
    fingerprints = {hashlib.sha256(ssl.PEM_cert_to_DER_cert(pem.decode("ascii"))).hexdigest() for pem in certificates}
    if not {ROOT_FINGERPRINT, SUB_FINGERPRINT} <= fingerprints:
        raise ValueError("Official certificate chain fingerprints mismatch")
    with urllib.request.urlopen(CURRENT_SUB, timeout=30) as response:
        sub_ca = response.read().decode("ascii")
    if hashlib.sha256(ssl.PEM_cert_to_DER_cert(sub_ca)).hexdigest() != CURRENT_SUB_FINGERPRINT:
        raise ValueError("Official intermediate certificate fingerprint mismatch")
    context = ssl.create_default_context()
    context.load_verify_locations(cadata=sub_ca)
    context.verify_flags |= ssl.VERIFY_X509_PARTIAL_CHAIN
    return context


def fetch_workbook(file_year, context):
    url = SOURCE_PATTERN.format(year=file_year)
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (benefit-calculator-review)"})
    try:
        with urllib.request.urlopen(request, timeout=60, context=context) as response:
            raw = response.read()
    except urllib.error.HTTPError as exc:
        if exc.code == 404:
            return None
        raise
    if len(raw) > 15_000_000 or not raw.startswith(b"PK"):
        raise ValueError(f"Unexpected Rosstat workbook: {url}")
    return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def merge_years(existing, found, preliminary, sources, checked):
    updated = json.loads(json.dumps(existing))
    changes = []
    for year, regions in sorted(found.items()):
        if len(regions) < 80:
            raise ValueError(f"Only {len(regions)} regions found for {year}; workbook layout needs review")
        if year == 2024 and abs(regions.get("Республика Мордовия", 0) - 57133.1) > 1:
            raise ValueError("2024 Mordovia control value failed")
        key = str(year)
        old = updated["years"].get(key, {})
        changed = old.get("regions") != regions or old.get("preliminary") != (year in preliminary)
        if not changed:
            continue
        # Any revision invalidates the old review until a person checks it.
        updated["years"][key] = {
            "regions": regions,
            "preliminary": year in preliminary,
            "final_confirmed": False,
            "publication_month": None,
            "source": sources[year],
            "first_seen": old.get("first_seen") or checked,
            "changed_at": checked,
        }
        changes.append(f"{year}: {len(regions)} регионов, {'предварительные' if year in preliminary else 'без пометки о предварительности'}; требует проверки публикации")
    return updated, changes


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--update", action="store_true")
    parser.add_argument("--year", type=int, default=dt.date.today().year)
    args = parser.parse_args()
    existing = json.loads(DATA.read_text(encoding="utf-8"))
    context = verified_rosstat_context()
    found, preliminary, sources = {}, set(), {}
    # The newest workbook may not exist in March. The previous one remains a
    # valid source for year - 2 while we continue to check the new year weekly.
    for file_year in [args.year - 1, args.year - 2]:
        if file_year < 2025:
            continue
        book = fetch_workbook(file_year, context)
        if book is None:
            print(f"Файл годовой таблицы за {file_year} год пока не опубликован.")
            continue
        years, flagged = extract(book)
        for year, regions in years.items():
            if year in [args.year - 1, args.year - 2] and year not in found:
                found[year], sources[year] = regions, SOURCE_PATTERN.format(year=file_year)
                if year in flagged:
                    preliminary.add(year)
    if not found:
        raise ValueError("Нет доступной годовой таблицы Росстата")
    updated, changes = merge_years(existing, found, preliminary, sources, dt.date.today().isoformat())
    for change in changes:
        print(change)
    if not changes:
        print("Годовые показатели не изменились.")
    if args.update and changes:
        DATA.write_text(json.dumps(updated, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
